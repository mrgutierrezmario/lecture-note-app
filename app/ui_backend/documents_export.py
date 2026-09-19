"""Lecture exports as PDF (reportlab) and Word (python-docx).

One document per lecture with up to three sections — the notes, the questions
asked in "Ask about the lecture" with their answers, and the full transcript
grouped into ~30-second paragraphs with a time column. Both formats share the
same structure and the app's restrained branding (logo mark, blue section
headings, page header/footer on the PDF).

The Markdown the notes are stored as is simple (``#``/``##`` headings, ``-``
bullets, ``**bold**``, ``*italic*``, `` `code` ``), so it is converted with a
small purpose-built parser rather than a Markdown library.
"""

import io
import re
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches as DocxInches
from docx.shared import Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    HRFlowable,
    KeepTogether,
    ListFlowable,
    ListItem,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

APP_NAME = "AI Lecture Notes"
BRAND = "M.G. NETWORK AND TECHNOLOGY SOLUTIONS"
# The logo mark ships with the frontend: under dist/ in the Docker image (the
# built site), under public/ in a development checkout.
_UI = Path(__file__).resolve().parent.parent / "ui_frontend"
LOGO = next(
    (
        p
        for p in (
            _UI / "dist" / "android-chrome-512x512.png",
            _UI / "public" / "android-chrome-512x512.png",
        )
        if p.exists()
    ),
    _UI / "public" / "android-chrome-512x512.png",
)
PARAGRAPH_SECONDS = 30  # transcript segments are grouped into paragraphs this long

# Palette (matches the UI's light theme)
BLUE = "#0b74f6"
INK = "#0f1f3d"
MUTED = "#5b6b86"
FAINT = "#8a97ae"
RULE = "#dfe5ef"
TINT = "#eef4ff"


@dataclass
class LectureDoc:
    """Everything an export needs, gathered by the route from the database."""

    title: str
    recorded_at: datetime | None
    duration_seconds: int
    notes_version: int
    notes_md: str | None  # None/empty = no notes section
    transcript: list[tuple[float, str]] = field(default_factory=list)  # (seconds, text)
    qa: list[tuple[str, str, str | None]] = field(default_factory=list)  # (role, text, provider)

    @property
    def subtitle(self) -> str:
        """ "Recorded …, · 1 h 53 min · Notes version 29"."""
        parts = []
        if self.recorded_at:
            parts.append(f"Recorded {self.recorded_at:%B %d, %Y, %I:%M %p}")
        m = self.duration_seconds // 60
        parts.append(f"{m // 60} h {m % 60} min" if m >= 60 else f"{m} min")
        if self.notes_version:
            parts.append(f"Notes version {self.notes_version}")
        return " · ".join(parts)

    def sections(self) -> list[str]:
        """Which of notes / qa / transcript have content, in document order."""
        out = []
        if self.notes_md and self.notes_md.strip():
            out.append("notes")
        if self.qa:
            out.append("qa")
        if self.transcript:
            out.append("transcript")
        return out


def group_transcript(segments: list[tuple[float, str]]) -> list[tuple[float, str]]:
    """Merge (seconds, text) segments into ~PARAGRAPH_SECONDS paragraphs."""
    out, buf, start = [], [], None
    for ts, text in segments:
        if start is None:
            start = ts
        if ts - start >= PARAGRAPH_SECONDS and buf:
            out.append((start, " ".join(buf)))
            buf, start = [], ts
        buf.append(text.strip())
    if buf:
        out.append((start or 0.0, " ".join(buf)))
    return out


def _mmss(seconds: float) -> str:
    return f"{int(seconds // 60):02d}:{int(seconds % 60):02d}"


# ── Markdown (the subset the notes use) ───────────────────────────────────────


def _md_blocks(md: str):
    """Yield ("h1"|"h2"|"bullet"|"p", text) for each non-empty Markdown line."""
    for line in md.splitlines():
        s = line.strip()
        if not s:
            continue
        if s.startswith("## "):
            yield "h2", s[3:]
        elif s.startswith("# "):
            yield "h1", s[2:]
        elif s.startswith(("- ", "* ")):
            yield "bullet", s[2:]
        else:
            yield "p", s


def _inline_runs(text: str) -> list[tuple[str, dict]]:
    """Split inline Markdown into (text, {bold, italic, code}) runs."""
    runs: list[tuple[str, dict]] = []
    pattern = re.compile(r"(\*\*.+?\*\*|(?<!\*)\*(?!\*).+?\*(?!\*)|`.+?`)")
    for piece in pattern.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            runs.append((piece[2:-2], {"bold": True}))
        elif piece.startswith("`") and piece.endswith("`"):
            runs.append((piece[1:-1], {"code": True}))
        elif piece.startswith("*") and piece.endswith("*"):
            runs.append((piece[1:-1], {"italic": True}))
        else:
            runs.append((piece, {}))
    return runs


def _rl_inline(text: str) -> str:
    """Inline Markdown → reportlab's mini-HTML."""
    text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    text = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", text)
    text = re.sub(r"(?<!\*)\*(?!\*)(.+?)\*(?!\*)", r"<i>\1</i>", text)
    text = re.sub(r"`(.+?)`", r"<font face='Courier'>\1</font>", text)
    return text


# ── PDF ───────────────────────────────────────────────────────────────────────

_ss = getSampleStyleSheet()
_BASE = ParagraphStyle(
    "base",
    parent=_ss["Normal"],
    fontName="Helvetica",
    fontSize=10,
    leading=14,
    textColor=colors.HexColor(INK),
)
_STYLES = {
    "title": ParagraphStyle(
        "title", parent=_BASE, fontName="Helvetica-Bold", fontSize=20, leading=24, spaceAfter=2
    ),
    "sub": ParagraphStyle(
        "sub", parent=_BASE, fontSize=9.5, textColor=colors.HexColor(MUTED), spaceAfter=10
    ),
    "h1": ParagraphStyle(
        "h1",
        parent=_BASE,
        fontName="Helvetica-Bold",
        fontSize=15,
        leading=19,
        spaceBefore=16,
        spaceAfter=6,
    ),
    "h2": ParagraphStyle(
        "h2",
        parent=_BASE,
        fontName="Helvetica-Bold",
        fontSize=11.5,
        leading=15,
        textColor=colors.HexColor(BLUE),
        spaceBefore=12,
        spaceAfter=4,
    ),
    "bullet": ParagraphStyle("bullet", parent=_BASE, spaceAfter=2),
    "ts": ParagraphStyle(
        "ts",
        parent=_BASE,
        fontName="Courier",
        fontSize=8.5,
        leading=13,
        textColor=colors.HexColor(FAINT),
    ),
    "txt": ParagraphStyle("txt", parent=_BASE, fontSize=9.5, leading=13.5),
    "q": ParagraphStyle("q", parent=_BASE, fontName="Helvetica-Bold", fontSize=10),
    "a": ParagraphStyle("a", parent=_BASE, fontSize=9.5, leading=13.5, spaceAfter=5),
    "label": ParagraphStyle(
        "label", parent=_BASE, fontSize=7.5, leading=10, textColor=colors.HexColor(FAINT)
    ),
}


def _pdf_markdown(md: str, body_style: str = "base") -> list:
    """Markdown → reportlab flowables (headings, bullet lists, paragraphs)."""
    out, items = [], []

    def flush():
        nonlocal items
        if items:
            out.append(
                ListFlowable(
                    [
                        ListItem(
                            Paragraph(_rl_inline(i), _STYLES["bullet"]), leftIndent=12, value="•"
                        )
                        for i in items
                    ],
                    bulletType="bullet",
                    start="•",
                    bulletFontSize=8,
                    bulletColor=colors.HexColor(BLUE),
                    leftIndent=12,
                )
            )
            items = []

    for kind, text in _md_blocks(md):
        if kind == "bullet":
            items.append(text)
            continue
        flush()
        out.append(
            Paragraph(
                _rl_inline(text),
                _STYLES["h2" if kind == "h2" else "h1" if kind == "h1" else body_style],
            )
        )
    flush()
    return out


def _pdf_page_frame(title: str, generated: str):
    def draw(canvas, doc):
        canvas.saveState()
        w, h = letter
        left, right = 0.75 * inch, w - 0.75 * inch
        if LOGO.exists():
            canvas.drawImage(
                str(LOGO), left, h - 0.72 * inch, width=0.32 * inch, height=0.32 * inch, mask="auto"
            )
        canvas.setFont("Helvetica-Bold", 9)
        canvas.setFillColor(colors.HexColor(INK))
        canvas.drawString(left + 0.42 * inch, h - 0.53 * inch, APP_NAME)
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor(FAINT))
        canvas.drawString(left + 0.42 * inch, h - 0.65 * inch, BRAND)
        canvas.setFont("Helvetica", 8.5)
        canvas.setFillColor(colors.HexColor(MUTED))
        canvas.drawRightString(right, h - 0.58 * inch, title[:70])
        canvas.setStrokeColor(colors.HexColor(RULE))
        canvas.setLineWidth(0.6)
        canvas.line(left, h - 0.82 * inch, right, h - 0.82 * inch)
        canvas.line(left, 0.7 * inch, right, 0.7 * inch)
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor(FAINT))
        canvas.drawString(left, 0.5 * inch, f"Generated by {APP_NAME} · {generated}")
        canvas.drawRightString(right, 0.5 * inch, f"Page {doc.page}")
        canvas.restoreState()

    return draw


def build_pdf(lec: LectureDoc) -> bytes:
    """Render the lecture document as PDF bytes."""
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=letter,
        leftMargin=0.75 * inch,
        rightMargin=0.75 * inch,
        topMargin=1.05 * inch,
        bottomMargin=0.95 * inch,
        title=lec.title,
        author=APP_NAME,
    )
    S = _STYLES
    sections = lec.sections()
    names = {"notes": "Notes", "qa": "Questions &amp; answers", "transcript": "Full transcript"}
    contents = " &nbsp;&nbsp;·&nbsp;&nbsp; ".join(
        f"{i}. {names[k]}" for i, k in enumerate(sections, 1)
    )

    def rule():
        return HRFlowable(width="100%", thickness=0.6, color=colors.HexColor(RULE), spaceAfter=4)

    story = [Paragraph(_rl_inline(lec.title), S["title"]), Paragraph(lec.subtitle, S["sub"])]
    if len(sections) > 1:
        story.append(
            Table(
                [[Paragraph(f"<b>Contents</b>&nbsp;&nbsp; {contents}", S["sub"])]],
                colWidths=[7 * inch],
                style=TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(TINT)),
                        ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor(RULE)),
                        ("LEFTPADDING", (0, 0), (-1, -1), 10),
                        ("TOPPADDING", (0, 0), (-1, -1), 6),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                    ]
                ),
            )
        )

    for i, key in enumerate(sections, 1):
        if i > 1:
            story.append(PageBreak())
        story += [Paragraph(f"{i}. {names[key]}", S["h1"]), rule()]
        if key == "notes":
            story += _pdf_markdown(lec.notes_md or "")
        elif key == "qa":
            story.append(
                Paragraph(
                    "Asked in “Ask about the lecture”. Answers are drawn only from this lecture's "
                    "transcript, notes and slides.",
                    S["sub"],
                )
            )
            for role, text, provider in lec.qa:
                if role == "user":
                    story.append(
                        KeepTogether(
                            [
                                Paragraph("YOU", S["label"]),
                                Paragraph(_rl_inline(text), S["q"]),
                                Spacer(1, 4),
                            ]
                        )
                    )
                else:
                    # Answers keep their paragraphs and bullets (the model's
                    # formatting), rendered inside the tinted block.
                    block = Table(
                        [[_pdf_markdown(text, body_style="a") or [Paragraph("", S["a"])]]],
                        colWidths=[6.8 * inch],
                        style=TableStyle(
                            [
                                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(TINT)),
                                ("LINEBEFORE", (0, 0), (0, -1), 2, colors.HexColor(BLUE)),
                                ("LEFTPADDING", (0, 0), (-1, -1), 10),
                                ("RIGHTPADDING", (0, 0), (-1, -1), 10),
                                ("TOPPADDING", (0, 0), (-1, -1), 7),
                                ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
                            ]
                        ),
                    )
                    label = f"AI · {_provider_label(provider)}" if provider else "AI"
                    story += [Paragraph(label, S["label"]), block, Spacer(1, 10)]
        else:
            story.append(
                Paragraph("Times are minutes:seconds from the start of the recording.", S["sub"])
            )
            rows = [
                [Paragraph(_mmss(ts), S["ts"]), Paragraph(_rl_inline(text), S["txt"])]
                for ts, text in group_transcript(lec.transcript)
            ]
            story.append(
                Table(
                    rows,
                    colWidths=[0.6 * inch, 6.4 * inch],
                    style=TableStyle(
                        [
                            ("VALIGN", (0, 0), (-1, -1), "TOP"),
                            ("TOPPADDING", (0, 0), (-1, -1), 4),
                            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                            ("LEFTPADDING", (0, 0), (-1, -1), 0),
                            ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                        ]
                    ),
                )
            )

    frame = _pdf_page_frame(lec.title, datetime.now().strftime("%B %d, %Y"))
    doc.build(story, onFirstPage=frame, onLaterPages=frame)
    return buf.getvalue()


def _provider_label(provider: str | None) -> str:
    if not provider:
        return ""
    head = provider.split("/")[0]
    return {
        "gemini": "Gemini",
        "claude": "Claude",
        "openai": "OpenAI",
        "ollama": "Local model",
        "llava": "Local vision model",
        "blip": "Local caption",
    }.get(head, head)


# ── Word ──────────────────────────────────────────────────────────────────────


def _docx_shade(paragraph, hex_fill: str) -> None:
    """Background colour on a paragraph (python-docx has no API for it)."""
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill.lstrip("#"))
    p_pr.append(shd)


def _docx_runs(paragraph, text: str, size: float | None = None, color: str | None = None) -> None:
    for piece, fmt in _inline_runs(text):
        run = paragraph.add_run(piece)
        run.bold = fmt.get("bold", False) or None
        run.italic = fmt.get("italic", False) or None
        if fmt.get("code"):
            run.font.name = "Consolas"
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = RGBColor.from_string(color.lstrip("#"))


def build_docx(lec: LectureDoc) -> bytes:
    """Render the lecture document as a .docx."""
    d = Document()
    for section in d.sections:
        section.left_margin = section.right_margin = DocxInches(0.9)
        section.top_margin = section.bottom_margin = DocxInches(0.9)
        header = section.header.paragraphs[0]
        if LOGO.exists():
            header.add_run().add_picture(str(LOGO), width=DocxInches(0.28))
            header.add_run("  ")
        r = header.add_run(APP_NAME)
        r.bold = True
        r.font.size = Pt(9)
        r2 = header.add_run(f"   ·   {BRAND.title()}   ·   {lec.title}")
        r2.font.size = Pt(8)
        r2.font.color.rgb = RGBColor.from_string(FAINT.lstrip("#"))
        footer = section.footer.paragraphs[0]
        fr = footer.add_run(f"Generated by {APP_NAME} · {datetime.now():%B %d, %Y}")
        fr.font.size = Pt(8)
        fr.font.color.rgb = RGBColor.from_string(FAINT.lstrip("#"))
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    normal = d.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    t = d.add_paragraph()
    _docx_runs(t, lec.title, size=20, color=INK)
    t.runs[0].bold = True
    sub = d.add_paragraph()
    _docx_runs(sub, lec.subtitle, size=9.5, color=MUTED)

    sections = lec.sections()
    names = {"notes": "Notes", "qa": "Questions & answers", "transcript": "Full transcript"}
    for i, key in enumerate(sections, 1):
        if i > 1:
            d.add_page_break()
        h = d.add_heading(f"{i}. {names[key]}", level=1)
        for run in h.runs:
            run.font.color.rgb = RGBColor.from_string(INK.lstrip("#"))
        if key == "notes":
            for kind, text in _md_blocks(lec.notes_md or ""):
                if kind in ("h1", "h2"):
                    hh = d.add_heading("", level=2 if kind == "h2" else 1)
                    _docx_runs(hh, text)
                    for run in hh.runs:
                        run.font.color.rgb = RGBColor.from_string(BLUE.lstrip("#"))
                elif kind == "bullet":
                    _docx_runs(d.add_paragraph(style="List Bullet"), text)
                else:
                    _docx_runs(d.add_paragraph(), text)
        elif key == "qa":
            note = d.add_paragraph()
            _docx_runs(
                note,
                "Asked in “Ask about the lecture”. Answers are drawn only from this "
                "lecture's transcript, notes and slides.",
                size=9.5,
                color=MUTED,
            )
            for role, text, provider in lec.qa:
                if role == "user":
                    lab = d.add_paragraph()
                    _docx_runs(lab, "YOU", size=7.5, color=FAINT)
                    lab.paragraph_format.space_after = Pt(0)
                    q = d.add_paragraph()
                    _docx_runs(q, text)
                    for run in q.runs:
                        run.bold = True
                else:
                    lab = d.add_paragraph()
                    _docx_runs(
                        lab,
                        f"AI · {_provider_label(provider)}" if provider else "AI",
                        size=7.5,
                        color=FAINT,
                    )
                    lab.paragraph_format.space_after = Pt(0)
                    blocks = list(_md_blocks(text)) or [("p", "")]
                    for n, (kind, block_text) in enumerate(blocks):
                        a = d.add_paragraph(style="List Bullet" if kind == "bullet" else None)
                        _docx_runs(a, block_text, size=10)
                        if kind in ("h1", "h2"):
                            for run in a.runs:
                                run.bold = True
                        a.paragraph_format.left_indent = DocxInches(
                            0.15 if kind != "bullet" else 0.4
                        )
                        _docx_shade(a, TINT)
                        a.paragraph_format.space_after = Pt(10 if n == len(blocks) - 1 else 2)
        else:
            note = d.add_paragraph()
            _docx_runs(
                note,
                "Times are minutes:seconds from the start of the recording.",
                size=9.5,
                color=MUTED,
            )
            table = d.add_table(rows=0, cols=2)
            table.autofit = True
            for ts, text in group_transcript(lec.transcript):
                cells = table.add_row().cells
                cells[0].width = DocxInches(0.7)
                cells[1].width = DocxInches(6.0)
                _docx_runs(cells[0].paragraphs[0], _mmss(ts), size=8.5, color=FAINT)
                cells[0].paragraphs[0].runs[0].font.name = "Consolas"
                _docx_runs(cells[1].paragraphs[0], text, size=10)

    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
